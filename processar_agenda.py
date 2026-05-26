import pandas as pd
from openpyxl.utils import get_column_letter

# ── Configurações ────────────────────────────────────────────────────────────

ARQUIVO_ENTRADA = "table_6a04a78892761.csv"
ARQUIVO_SAIDA   = "agenda_processada.xlsx"

DATA_REAJUSTE_PARTICULAR = pd.Timestamp("2025-10-01")

COLUNAS_REMOVER = {
    "CPF", "Criado por", "Criado Por", "Sala", "Serviço",
    "Pagamento", "Comissão Total", "Comissão Paga", "Total de Comissão", "Etiqueta",
}

PROFISSIONAL_EXCLUIR = "ana maria faria dos santos"

# ── Utilitários ──────────────────────────────────────────────────────────────

def formatar_telefone(val):
    """Normaliza números de telefone: inteiros ganham '+', strings inválidas viram ''."""
    if pd.isna(val):
        return ""
    if isinstance(val, (int, float)):
        return "+" + str(int(val))
    texto = str(val).strip()
    return "" if texto.lower() in ("nan", "none") else texto


def calcular_idade_na_consulta(nascimento: pd.Series, data_consulta: pd.Series) -> pd.Series:
    """Calcula idade completa na data da consulta, respeitando se o aniversário já passou."""
    validos = nascimento.notna() & data_consulta.notna()
    aniversario_futuro = (data_consulta.dt.month < nascimento.dt.month) | (
        (data_consulta.dt.month == nascimento.dt.month)
        & (data_consulta.dt.day < nascimento.dt.day)
    )
    return (
        (data_consulta.dt.year - nascimento.dt.year - aniversario_futuro.astype(int))
        .where(validos)
    )


def salvar_excel(df: pd.DataFrame, caminho: str):
    """Salva o DataFrame em .xlsx, garantindo que 'Telefone' seja tratado como texto."""
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Dados")
        if "Telefone" not in df.columns:
            return
        ws = writer.sheets["Dados"]
        col_letra = get_column_letter(list(df.columns).index("Telefone") + 1)
        for linha in range(2, len(df) + 2):
            celula = ws[f"{col_letra}{linha}"]
            celula.number_format = "@"
            celula.value = str(celula.value).strip() if celula.value else ""


# ── Pipeline principal ───────────────────────────────────────────────────────

# 1. Leitura do CSV
df = pd.read_csv(ARQUIVO_ENTRADA, encoding="utf-8")

# 2. Formatar telefone como texto
if "Telefone" in df.columns:
    df["Telefone"] = df["Telefone"].map(formatar_telefone)

# 3. Converter datas
df["Data"]               = pd.to_datetime(df["Data"],               dayfirst=True, errors="coerce")
df["Data de Nascimento"] = pd.to_datetime(df["Data de Nascimento"], dayfirst=True, errors="coerce")

# 4. Remover profissional excluído
df = df[df["Profissional"].astype(str).str.strip().str.lower() != PROFISSIONAL_EXCLUIR]

# 5. Normalizar status ("Confirmado" → "Atendido")
if "Status" in df.columns:
    df["Status"] = df["Status"].astype(str).str.strip().replace({"Confirmado": "Atendido"})

# 6. Calcular idade na data da consulta
df["Idade na Consulta"] = calcular_idade_na_consulta(df["Data de Nascimento"], df["Data"])

# 7. Derivar flags de negócio
tipo          = df["Convênio"].fillna("").astype(str).str.strip().str.lower()
eh_particular = tipo.str.contains("particular")
eh_amil       = tipo.str.contains("amil") & ~eh_particular
apos_reajuste = df["Data"] >= DATA_REAJUSTE_PARTICULAR
menor_idade   = df["Idade na Consulta"] <= 18

# 8. Aplicar valores e comissões
valor    = pd.Series(0.0, index=df.index)
comissao = pd.Series(0.0, index=df.index)

valor.loc[eh_particular &  apos_reajuste] = 110
valor.loc[eh_particular & ~apos_reajuste] = 95
valor.loc[eh_amil       &  menor_idade]   = 80
valor.loc[eh_amil       & ~menor_idade]   = 50

comissao.loc[eh_particular &  apos_reajuste] = 55
comissao.loc[eh_particular & ~apos_reajuste] = 35
comissao.loc[eh_amil]                        = 25

df["Total"]    = valor
df["Comissão"] = comissao

# 9. Remover colunas desnecessárias e exportar
df.drop(columns=COLUNAS_REMOVER, inplace=True, errors="ignore")
salvar_excel(df, ARQUIVO_SAIDA)

print(f"Planilha salva em: {ARQUIVO_SAIDA}")
print(f"Linhas processadas: {len(df)}")